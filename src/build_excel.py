"""Рабочая книга Excel с формулами методики сравнительной оценки ЭТ и Д.

Все расчёты — формулами Excel от именованных исходных данных; опытные данные — значения из журналов.
Результаты Монте-Карло и зоны (требуют многократных пересчётов) приведены значениями из Python.
Запуск: python3 src/build_excel.py
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName

from models import AHP_GROUPS, AHP_LOCAL, CRITERIA, GROUPS
from params import P, RANGES

ROOT = Path(__file__).resolve().parents[1]
T = ROOT / "results" / "tables"
OUT = ROOT / "excel" / "Расчёты_сравнение_ЭТ_Д.xlsx"

FONT = "Times New Roman"
F_IN = Font(name=FONT, color="0000FF")
F_LINK = Font(name=FONT, color="008000")
F_N = Font(name=FONT)
F_B = Font(name=FONT, bold=True)
F_H = Font(name=FONT, bold=True, color="FFFFFF")
FILL_H = PatternFill("solid", fgColor="2A78D6")
FILL_IN = PatternFill("solid", fgColor="FFF2CC")
FILL_KEY = PatternFill("solid", fgColor="FFFF00")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CODES = [c[0] for c in CRITERIA]

wb = Workbook()


def name(n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


def header(ws, row, cols, widths=None):
    for j, h in enumerate(cols, 1):
        c = ws.cell(row, j, h)
        c.font, c.fill, c.border = F_H, FILL_H, BOX
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[L(j)].width = w
    ws.row_dimensions[row].height = 45


def put(ws, r, c, v, font=F_N, fmt=None, fill=None):
    cell = ws.cell(r, c, v)
    cell.font, cell.border = font, BOX
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    return cell


# ============================================================ Описание
ws = wb.active
ws.title = "Описание"
ws.column_dimensions["A"].width = 28
ws.column_dimensions["B"].width = 100
lines = [
    ("Рабочая книга", "Сравнительная оценка электротрактора JINMA YL254ET (ЭТ) и дизельного аналога JINMA JM-254 (Д) по результатам полевых испытаний 15.09–15.10.2026"),
    ("Назначение", "Расчёт показателей обеих машин, нормированных преимуществ, весов (МАИ + энтропийный метод) и интегрального индекса I для 16 исследованных режимов; пересчёт при изменении исходных данных"),
    ("Как пользоваться", "Изменяйте только ячейки с синим шрифтом на жёлтом фоне на листах «Исходные» и «Веса». Все остальные значения пересчитываются формулами"),
    ("Ключевые рычаги", "S_task (сменное задание; 0 — полная смена), T_year_h (годовая загрузка), c_D и c_el (цены), l_charge_km (удалённость зарядки), P_ch_ac (мощность зарядки), цены машин"),
    ("Цвета", "синий шрифт — исходные данные; чёрный — формулы; зелёный — ссылки на другие листы; жёлтая заливка — изменяемые ячейки"),
    ("Лист «Исходные»", "исходные данные и допущения с источниками и диапазонами вариации (именованные ячейки)"),
    ("Лист «Проходы»", "результаты зачётных проходов (значения из журналов CAN 1 Гц и полевых журналов) и расчётные показатели проходов (формулы)"),
    ("Лист «Опыт_режимы»", "средние по повторностям, стандартные отклонения, доверительные интервалы (формулы AVERAGEIFS и т. п.)"),
    ("Лист «Модели»", "коэффициенты моделей: сопротивление качению, k = k0 + ε·v² (ЛИНЕЙН), A = A0 + Av·v + Aa·a (ЛИНЕЙН), тяговая характеристика"),
    ("Лист «Энергетика»", "мощности ЭТ и Д, буксование, резерв мощности, допустимость режимов, часовой расход топлива Д (ASAE D497.7)"),
    ("Лист «Смена»", "модель смены: развороты, автономность, подзарядки (до трёх циклов), сменная выработка, энергия из сети, расход топлива на 1 га"),
    ("Лист «Экономика»", "удельные прямые и совокупные затраты (ГОСТ 34393-2018, ASAE D497.7)"),
    ("Лист «Показатели»", "показатели I1–I10 обеих машин, нормированные преимущества d и интегральный индекс I"),
    ("Лист «Веса»", "матрицы парных сравнений МАИ (изменяемые), веса, согласованность; итоговые веса"),
    ("Лист «Энтропия»", "энтропийные веса по матрице показателей режимов, реализуемых обеими машинами"),
    ("Лист «Устойчивость»", "результаты Монте-Карло (1000 испытаний) и чувствительности — значения из программного комплекса Python (требуют многократного пересчёта)"),
    ("Проверка", "при исходных данных по умолчанию значения индекса совпадают с отчётом (таблица 5.6, сценарий 1); расхождение собственного вектора МАИ и метода геометрического среднего < 0,001"),
]
ws["A1"] = "РАСЧЁТЫ К ОТЧЁТУ О НИР"
ws["A1"].font = Font(name=FONT, bold=True, size=14)
for i, (a, b) in enumerate(lines, 3):
    ws.cell(i, 1, a).font = F_B
    c = ws.cell(i, 2, b)
    c.font = F_N
    c.alignment = Alignment(wrap_text=True)

# ============================================================ Исходные
ws = wb.create_sheet("Исходные")
header(ws, 1, ["Обозначение (имя)", "Параметр", "Значение", "Ед.", "Статус, источник", "Мин.", "Макс."], [16, 52, 14, 12, 52, 10, 10])
SPEC = [
    ("m_ET", "Масса ЭТ", "кг", "паспорт изготовителя"), ("m_D", "Масса Д", "кг", "TractorData.com"),
    ("g", "Ускорение свободного падения", "м/с²", "константа"),
    ("E_bat_kWh", "Энергоёмкость батареи", "кВт·ч", "304 А·ч × 76,8 В; подтверждено опытом"),
    ("SOC_max", "SOC после зарядки", "%", "опыт"), ("SOC_min", "Нижняя граница SOC", "%", "порог предупреждения BMS (код 61)"),
    ("P_el_lim", "Предельная мощность от батареи", "кВт", "опыт (17,3–17,5 кВт)"),
    ("P_mot_1h", "Часовая мощность двигателя хода", "кВт", "паспорт"),
    ("P_aux", "Собственное потребление систем", "кВт", "опыт"),
    ("eta_ep", "КПД электропривода", "–", "допущение"), ("eta_tr", "КПД трансмиссии", "–", "допущение"),
    ("P_hyd_idle_D", "Потери гидронасоса Д на холостом ходу", "кВт", "допущение"),
    ("P_e_nom_D", "Номинальная мощность двигателя Д", "кВт", "паспорт"),
    ("k_pto_D", "Отношение P_ВОМ/P_е для Д", "–", "допущение"),
    ("k_D", "Поправка расхода топлива к ASAE D497.7", "–", "допущение (1,0 — в пользу Д)"),
    ("sfc_a", "ASAE D497.7: коэффициент a", "л/(кВт·ч)", "ASAE D497.7"), ("sfc_b", "ASAE D497.7: коэффициент b", "л/(кВт·ч)", "ASAE D497.7"),
    ("LHV_D", "Теплота сгорания ДТ", "МДж/л", "справочное"), ("tank_D", "Топливный бак Д", "л", "TractorData.com"),
    ("t_refuel_h", "Время дозаправки", "ч", "допущение"),
    ("P_ch_ac", "Мощность зарядного устройства", "кВт", "допущение: разъём 16 А, 230 В"),
    ("eta_ch", "КПД зарядного устройства", "–", "допущение"), ("eta_bat", "КПД батареи (энергетический)", "–", "допущение"),
    ("T_shift_h", "Продолжительность смены", "ч", "принято"), ("tau_svc", "Доля смены на работу и переезды", "–", "допущение"),
    ("L_run_m", "Длина гона", "м", "принято"), ("l_charge_km", "Удалённость места зарядки", "км", "принято"),
    ("v_transp_kmh", "Скорость переезда", "км/ч", "опыт"), ("P_transp_kW", "Мощность при переезде", "кВт", "опыт"),
    ("usd_kzt", "Курс доллара США", "тг", "НБРК, 18.10.2026"),
    ("price_ET_usd", "Цена ЭТ (FOB)", "$", "Made-in-China, Jiangsu Yueda"), ("price_D_usd", "Цена Д (FOB)", "$", "Made-in-China, Jinma"),
    ("k_landed", "Коэффициент «FOB → цена у дилера»", "–", "1,10 × 1,16 (НДС) × 1,15"),
    ("c_el", "Цена электроэнергии с НДС", "тг/(кВт·ч)", "ЭнергоСбыт АЖК с 12.09.2026"),
    ("c_D", "Цена дизельного топлива", "тг/л", "средняя розничная, сент.–окт. 2026"),
    ("wage_month", "Средняя зарплата в сельском хозяйстве", "тг/мес", "БНС АСПиР РК, II кв. 2026"),
    ("h_month", "Рабочих часов в месяце", "ч", "принято"), ("k_soc", "Коэффициент начислений на ФОТ", "–", "допущение"),
    ("T_year_h", "Годовая загрузка", "ч/год", "допущение"), ("life_y", "Срок службы", "лет", "допущение"),
    ("r_disc", "Реальная ставка дисконтирования", "–", "допущение"), ("salvage", "Остаточная стоимость", "доля", "допущение"),
    ("ASAE_RF1", "ASAE D497.7: RF1 (колёсные тракторы)", "–", "ASAE D497.7"), ("ASAE_RF2", "ASAE D497.7: RF2", "–", "ASAE D497.7"),
    ("k_RM_ET", "Доля затрат на ремонт и ТО ЭТ от Д", "–", "допущение"),
    ("bat_usd_kWh", "Стоимость замены батареи", "$/(кВт·ч)", "2,5 × BNEF-2025 LFP"),
    ("bat_cycles", "Цикловой ресурс батареи", "циклов", "допущение"), ("bat_cal_y", "Календарный ресурс батареи", "лет", "допущение"),
    ("EF_grid", "Коэффициент эмиссии энергосистемы РК", "кг/(кВт·ч)", "IEA-2025 (через ABB)"),
    ("EF_D", "Коэффициент эмиссии ДТ", "кг/л", "IPCC 2006"),
    ("PEF_el", "Коэффициент первичной энергии электроэнергии", "–", "Директива 2012/27/ЕС"),
    ("PEF_D", "Коэффициент первичной энергии ДТ", "–", "JEC WTW v5"),
    ("slip_max", "Допустимое буксование", "–", "допущение"),
    ("S_task", "Сменное задание (0 — полная смена: задание = выработке Д)", "га/смену", "сценарий"),
    ("alpha_w", "Доля экспертных весов в итоговых", "–", "методика (0,5)"),
]
src_key = {"ASAE_RF1": "RF1", "ASAE_RF2": "RF2"}
for i, (n, txt, unit, src) in enumerate(SPEC, 2):
    v = {"S_task": 0.0, "alpha_w": 0.5}.get(n, P.get(src_key.get(n, n)))
    put(ws, i, 1, n, F_B)
    put(ws, i, 2, txt)
    put(ws, i, 3, v, F_IN, fill=FILL_IN)
    put(ws, i, 4, unit)
    put(ws, i, 5, src)
    rng = RANGES.get(n)
    put(ws, i, 6, rng[0] if rng else None)
    put(ws, i, 7, rng[1] if rng else None)
    name(n, f"Исходные!$C${i}")
r0 = len(SPEC) + 3
ws.cell(r0, 1, "Производные величины").font = F_B
DER = [
    ("E_use", "Используемый запас энергии, кВт·ч", "=E_bat_kWh*(SOC_max-SOC_min)/100"),
    ("T_av", "Время смены на работу и переезды, ч", "=T_shift_h*tau_svc"),
    ("t_tr", "Время переезда в одну сторону, ч", "=l_charge_km/v_transp_kmh"),
    ("E_tr", "Энергия переезда ЭТ в одну сторону, кВт·ч", "=P_transp_kW*t_tr"),
    ("t_ch", "Время полной подзарядки, ч", "=E_use/(eta_ch*eta_bat*P_ch_ac)"),
    ("P_pto_nom", "Номинальная мощность ВОМ Д, кВт", "=k_pto_D*P_e_nom_D"),
    ("wage_h", "Стоимость часа работы тракториста, тг/ч", "=wage_month/h_month*k_soc"),
    ("price_ET", "Цена ЭТ у дилера, тг", "=price_ET_usd*usd_kzt*k_landed"),
    ("price_D", "Цена Д у дилера, тг", "=price_D_usd*usd_kzt*k_landed"),
    ("T_life", "Наработка за срок службы, ч", "=T_year_h*life_y"),
    ("rm_h_D", "Средние затраты на ремонт и ТО Д, тг/ч", "=ASAE_RF1*price_D*(T_life/1000)^ASAE_RF2/T_life"),
    ("k_CRF", "Коэффициент возврата капитала", "=IF(r_disc=0,1/life_y,r_disc*(1+r_disc)^life_y/((1+r_disc)^life_y-1))"),
    ("k_SFF", "Коэффициент фонда возмещения", "=IF(r_disc=0,1/life_y,r_disc/((1+r_disc)^life_y-1))"),
    ("C_bat", "Стоимость замены батареи, тг", "=bat_usd_kWh*E_bat_kWh*usd_kzt"),
    ("Q_turn", "Расход топлива Д на развороте, л/ч", "=k_D*(sfc_a*(eta_ep*(turn_P-P_aux)+P_hyd_idle_D)+sfc_b*P_pto_nom)"),
    ("Q_tr", "Расход топлива Д на переезде, л/ч", "=k_D*(sfc_a*(eta_ep*(P_transp_kW-P_aux)*m_D/m_ET+P_hyd_idle_D)+sfc_b*P_pto_nom)"),
]
for i, (n, txt, fml) in enumerate(DER, r0 + 1):
    put(ws, i, 1, n, F_B)
    put(ws, i, 2, txt)
    put(ws, i, 3, fml, fmt="#,##0.000")
    name(n, f"Исходные!$C${i}")
ws.freeze_panes = "A2"

# ============================================================ Проходы
ps = pd.read_csv(T / "passes.csv")
ps = ps[~ps.limit_attempt]
ws = wb.create_sheet("Проходы")
cols = ["Машина", "Сессия", "Проход", "Блок", "Вид", "a_пл, м", "v_пл, км/ч", "Диапазон", "n, с",
        "P̄, кВт", "s_P, кВт", "v_т, км/ч", "v, км/ч", "a, м", "s_a, м", "B, м", "n_ВОМ, мин⁻¹",
        "W0, га/ч", "e, кВт·ч/га", "δ", "e_V, кДж/м³", "v², м²/с²", "F_к, Н", "R, Н", "k, кН/м²", "φ_н", "A, кДж/м³"]
header(ws, 1, cols, [12, 7, 7, 6, 11] + [9] * 22)
order = []
for impl in ("культиватор", "фреза"):
    for kind in ("холостой", "рабочий"):
        order.append(ps[(ps.implement == impl) & (ps.kind == kind)])
rows = pd.concat(order)
blocks = {}
r = 2
for x in rows.itertuples():
    key = (x.implement, x.kind)
    blocks.setdefault(key, [r, r])[1] = r
    vals = [x.implement, x.session, int(x.pass_no), int(x.block) if x.block == x.block else None, x.kind,
            x.a_plan_m, x.v_plan_kmh, x.gear, int(x.n_s), round(x.power_kW, 4), round(x.power_sd_kW, 4),
            round(x.v_wheel_kmh, 4), x.v_fact_kmh, x.a_meas_mean_m if x.a_meas_mean_m == x.a_meas_mean_m else None,
            x.a_meas_sd_m if x.a_meas_sd_m == x.a_meas_sd_m else None, x.B_meas_m if x.B_meas_m == x.B_meas_m else None,
            round(x.pto_rpm, 2)]
    for j, v in enumerate(vals, 1):
        put(ws, r, j, v, F_IN if j >= 10 else F_N)
    work = x.kind == "рабочий"
    put(ws, r, 18, f"=IF(P{r}=\"\",\"\",0.1*P{r}*M{r})" if work else None, fmt="0.000")
    put(ws, r, 19, f"=IF(R{r}=\"\",\"\",J{r}/R{r})" if work else None, fmt="0.00")
    put(ws, r, 20, f"=1-M{r}/L{r}", fmt="0.000")
    put(ws, r, 21, f"=IF(R{r}=\"\",\"\",3.6*J{r}/(P{r}*N{r}*M{r}))" if work else None, fmt="0.0")
    put(ws, r, 22, f"=(M{r}/3.6)^2", fmt="0.000")
    if x.implement == "культиватор":
        put(ws, r, 23, f"=eta_ep*eta_tr*(J{r}-P_aux)*1000/(L{r}/3.6)", fmt="0")
        if work:
            put(ws, r, 24, f"=W{r}-f_roll*m_ET*g", fmt="0")
            put(ws, r, 25, f"=X{r}/(P{r}*N{r})/1000", fmt="0.00")
            put(ws, r, 26, f"=X{r}/(m_ET*g)", fmt="0.000")
    elif work:
        put(ws, r, 27, f"=(J{r}-P_aux-c_move*L{r}/3.6-P_rot0)/(P{r}*N{r}*M{r}/3.6)", fmt="0.0")
    r += 1
ws.freeze_panes = "A2"
cs0, cs1 = blocks[("культиватор", "холостой")]
cw0, cw1 = blocks[("культиватор", "рабочий")]
ts0, ts1 = blocks[("фреза", "холостой")]
tw0, tw1 = blocks[("фреза", "рабочий")]

# ============================================================ Модели
ws = wb.create_sheet("Модели")
header(ws, 1, ["Имя", "Коэффициент", "Значение", "Ед.", "Формула / способ определения"], [12, 46, 14, 14, 70])
bg = pd.read_csv(T / "background.csv")
tr = pd.read_csv(T / "turns.csv")
import json
S = json.loads((ROOT / "results" / "summary.json").read_text())
M = S["models"]
MOD = [
    ("f_roll", "Коэффициент сопротивления качению", f"=AVERAGE(Проходы!W{cs0}:W{cs1})/(m_ET*g)", "–", "по самопередвижению с культиватором (F_к/G)"),
    ("k_0", "k0 – удельное сопротивление при v → 0", f"=INDEX(LINEST(Проходы!Y{cw0}:Y{cw1},Проходы!V{cw0}:V{cw1}),1,2)*1000", "Н/м²", "ЛИНЕЙН(k; v²)"),
    ("eps", "ε – скоростной коэффициент", f"=INDEX(LINEST(Проходы!Y{cw0}:Y{cw1},Проходы!V{cw0}:V{cw1}),1,1)*1000", "Н·с²/м⁴", "ЛИНЕЙН(k; v²)"),
    ("R2_k", "R² модели k", f"=INDEX(LINEST(Проходы!Y{cw0}:Y{cw1},Проходы!V{cw0}:V{cw1},TRUE,TRUE),3,1)", "–", "ЛИНЕЙН, статистика"),
    ("delta_0", "δ0 – буксование при самопередвижении", f"=AVERAGE(Проходы!T{cs0}:T{cs1})", "–", "среднее по самопередвижению"),
    ("phi_max", "φ_max тяговой характеристики", round(M["phi_max"], 6), "–", "нелинейный МНК (Python, scipy.curve_fit); значение"),
    ("k_slip", "k_δ тяговой характеристики", round(M["k_slip"], 6), "–", "нелинейный МНК (Python); значение"),
    ("c_move", "Удельная мощность передвижения с фрезой", f"=AVERAGE(Проходы!AB{ts0}:AB{ts1})", "кВт·с/м", "(P − P_aux)/v_т по самопередвижению на П2"),
    ("P_rot0", "Мощность холостого вращения ротора сверх P_aux", round(S["P_rot0_el"], 4), "кВт", "опыт: холостое вращение ротора, ВОМ 540"),
    ("A_0", "A0", f"=INDEX(LINEST(Проходы!AA{tw0}:AA{tw1},Проходы!M{tw0}:N{tw1}),1,3)", "кДж/м³", "ЛИНЕЙН(A; v, a)"),
    ("A_v", "Av", f"=INDEX(LINEST(Проходы!AA{tw0}:AA{tw1},Проходы!M{tw0}:N{tw1}),1,2)", "кДж/(м³·км/ч)", "ЛИНЕЙН(A; v, a)"),
    ("A_a", "Aa", f"=INDEX(LINEST(Проходы!AA{tw0}:AA{tw1},Проходы!M{tw0}:N{tw1}),1,1)", "кДж/(м³·м)", "ЛИНЕЙН(A; v, a)"),
    ("slip_till", "Буксование с фрезой (среднее)", f"=AVERAGE(Проходы!T{tw0}:T{tw1})", "–", "среднее по рабочим проходам"),
    ("cv_P", "Коэффициент вариации мощности (медиана)", round(M["cv_P"], 6), "–", "медиана s_P/P̄ по рабочим проходам"),
    ("turn_t", "Продолжительность разворота", round(S["turn"]["t_s"], 3), "с", "опыт"),
    ("turn_P", "Мощность на развороте", round(S["turn"]["P_kW"], 4), "кВт", "опыт"),
]
for i, (n, txt, v, u, how) in enumerate(MOD, 2):
    put(ws, i, 1, n, F_B)
    put(ws, i, 2, txt)
    put(ws, i, 3, v, F_N if isinstance(v, str) else F_IN, fmt="0.0000")
    put(ws, i, 4, u)
    put(ws, i, 5, how)
    name(n, f"Модели!$C${i}")
# столбец AB в «Проходах» для c_move: (P − P_aux)/(v_т/3,6) самопередвижения фрезы
wsP = wb["Проходы"]
put(wsP, 1, 28, "(P−P_aux)/v_т, кВт·с/м", F_H, fill=FILL_H)
for rr in range(ts0, ts1 + 1):
    put(wsP, rr, 28, f"=(J{rr}-P_aux)/(L{rr}/3.6)", fmt="0.000")
# столбец A (уд. энергия) для фрезы — колонка AA

# ============================================================ Опыт_режимы
reg = pd.read_csv(T / "regimes_measured.csv")
ws = wb.create_sheet("Опыт_режимы")
header(ws, 1, ["Машина", "a, м", "v, км/ч", "n", "P̄, кВт", "s_P", "Δ0,95 P", "e, кВт·ч/га", "s_e", "Δ0,95 e", "δ", "W0, га/ч"], [12] + [10] * 11)
rng = f"Проходы!$A${cw0}:$A${tw1}"
for i, x in enumerate(reg.itertuples(), 2):
    put(ws, i, 1, x.implement)
    put(ws, i, 2, x.a_plan)
    put(ws, i, 3, x.v_plan)
    crit = f"{rng},$A{i},Проходы!$F${cw0}:$F${tw1},$B{i},Проходы!$G${cw0}:$G${tw1},$C{i},Проходы!$E${cw0}:$E${tw1},\"рабочий\""
    put(ws, i, 4, f"=COUNTIFS({crit})")
    for col, src, sd in ((5, "J", 6), (8, "S", 9)):
        put(ws, i, col, f"=AVERAGEIFS(Проходы!${src}${cw0}:${src}${tw1},{crit})", fmt="0.00")
        # стандартное отклонение по трём повторностям
        put(ws, i, sd, f"=SQRT(SUMPRODUCT((Проходы!$A${cw0}:$A${tw1}=$A{i})*(Проходы!$F${cw0}:$F${tw1}=$B{i})*(Проходы!$G${cw0}:$G${tw1}=$C{i})*(Проходы!$E${cw0}:$E${tw1}=\"рабочий\")*(Проходы!${src}${cw0}:${src}${tw1}-{L(col)}{i})^2)/(D{i}-1))", fmt="0.00")
        put(ws, i, sd + 1, f"=TINV(0.05,D{i}-1)*{L(sd)}{i}/SQRT(D{i})", fmt="0.00")
    put(ws, i, 11, f"=AVERAGEIFS(Проходы!$T${cw0}:$T${tw1},{crit})", fmt="0.000")
    put(ws, i, 12, f"=AVERAGEIFS(Проходы!$R${cw0}:$R${tw1},{crit})", fmt="0.000")

# ============================================================ Энергетика
NR = len(reg)
ws = wb.create_sheet("Энергетика")
E_COLS = ["Машина", "a, м", "v, км/ч", "B, м", "v, м/с", "A, кДж/м³ (фреза)", "R, Н", "φ_н ЭТ", "δ ЭТ", "φ_н Д", "δ Д",
          "P_эл ЭТ, кВт", "P_вх ЭТ, кВт", "P_вх Д, кВт", "Резерв ЭТ", "Резерв Д", "Допустим ЭТ", "Допустим Д",
          "Длительный ЭТ", "Q Д, л/ч", "X Д"]
header(ws, 1, E_COLS, [12] + [10] * 20)
for i in range(2, NR + 2):
    k = i
    put(ws, i, 1, f"=Опыт_режимы!A{k}", F_LINK)
    put(ws, i, 2, f"=Опыт_режимы!B{k}", F_LINK)
    put(ws, i, 3, f"=Опыт_режимы!C{k}", F_LINK)
    cult = f'A{i}="культиватор"'
    put(ws, i, 4, f"=IF({cult},1.2,1.25)")
    put(ws, i, 5, f"=C{i}/3.6", fmt="0.000")
    put(ws, i, 6, f"=IF({cult},\"\",A_0+A_v*C{i}+A_a*B{i})", fmt="0.0")
    put(ws, i, 7, f"=IF({cult},D{i}*B{i}*(k_0+eps*E{i}^2),\"\")", fmt="0")
    put(ws, i, 8, f"=IF({cult},G{i}/(m_ET*g),\"\")", fmt="0.000")
    put(ws, i, 9, f"=IF({cult},delta_0-LN(1-H{i}/phi_max)/k_slip,slip_till)", fmt="0.000")
    put(ws, i, 10, f"=IF({cult},G{i}/(m_D*g),\"\")", fmt="0.000")
    put(ws, i, 11, f"=IF({cult},delta_0-LN(1-J{i}/phi_max)/k_slip,slip_till)", fmt="0.000")
    put(ws, i, 12, f"=IF({cult},P_aux+M{i}/eta_ep,P_aux+c_move*E{i}/(1-slip_till)+P_rot0+F{i}*D{i}*B{i}*E{i})", fmt="0.00")
    put(ws, i, 13, f"=IF({cult},(G{i}+f_roll*m_ET*g)*E{i}/(1-I{i})/eta_tr/1000,eta_ep*(L{i}-P_aux))", fmt="0.00")
    put(ws, i, 14, f"=IF({cult},(G{i}+f_roll*m_D*g)*E{i}/(1-K{i})/eta_tr/1000,eta_ep*(c_move*E{i}/(1-slip_till)*m_D/m_ET+P_rot0+F{i}*D{i}*B{i}*E{i}))+P_hyd_idle_D", fmt="0.00")
    put(ws, i, 15, f"=1-L{i}*(1+2*cv_P)/P_el_lim", fmt="0.0%")
    put(ws, i, 16, f"=1-N{i}*(1+2*cv_P)/P_pto_nom", fmt="0.0%")
    put(ws, i, 17, f"=AND(O{i}>=0,I{i}<=slip_max)")
    put(ws, i, 18, f"=AND(P{i}>=0,K{i}<=slip_max)")
    put(ws, i, 19, f"=AND(Q{i},M{i}<=P_mot_1h)")
    put(ws, i, 20, f"=k_D*(sfc_a*N{i}+sfc_b*P_pto_nom)", fmt="0.00")
    put(ws, i, 21, f"=N{i}/P_pto_nom", fmt="0.00")
ws.freeze_panes = "B2"

# ============================================================ Смена
ws = wb.create_sheet("Смена")
SH = ["Режим", "W0, га/ч", "n_пов, 1/га", "t_га, ч/га", "e_акб, кВт·ч/га", "Мощность цикла, кВт", "t_авт ЭТ, ч", "A_зар, га",
      "Работа на зарядке, ч", "w1, ч", "Подзаряд. 1", "w2, ч", "Подзаряд. 2", "w3, ч", "Подзарядок", "S_см ЭТ, га",
      "e_сеть, кВт·ч/га", "q0, л/га", "Q_ср, л/ч", "t_авт Д, ч", "Дозаправок", "S_см Д, га", "q, л/га", "A_бак, га"]
header(ws, 1, SH, [16] + [10] * 23)
for i in range(2, NR + 2):
    E = f"Энергетика!"
    put(ws, i, 1, f'=Энергетика!A{i}&" "&Энергетика!B{i}&"×"&Энергетика!C{i}', F_LINK)
    put(ws, i, 2, f"=0.1*{E}D{i}*{E}C{i}", fmt="0.000")
    put(ws, i, 3, f"=10000/({E}D{i}*L_run_m)", fmt="0.00")
    put(ws, i, 4, f"=1/B{i}+C{i}*turn_t/3600", fmt="0.000")
    put(ws, i, 5, f"={E}L{i}/B{i}+C{i}*turn_P*turn_t/3600", fmt="0.00")
    put(ws, i, 6, f"=E{i}/D{i}", fmt="0.00")
    put(ws, i, 7, f"=E_use/F{i}", fmt="0.00")
    put(ws, i, 8, f"=G{i}/D{i}", fmt="0.000")
    put(ws, i, 9, f"=(E_use-2*E_tr)/F{i}", fmt="0.00")
    # цикл 1
    put(ws, i, 10, f"=MAX(0,MIN(I{i},T_av-2*t_tr))", fmt="0.00")
    put(ws, i, 11, f"=AND(T_av-(t_tr+J{i})-t_tr>0.000001,T_av-(t_tr+J{i})-2*t_tr-t_ch>=0.25)")
    put(ws, i, 12, f"=IF(K{i},MAX(0,MIN(I{i},T_av-(t_tr+J{i}+2*t_tr+t_ch)-t_tr)),0)", fmt="0.00")
    t2e = f"(t_tr+J{i}+2*t_tr+t_ch+L{i})"
    put(ws, i, 13, f"=AND(K{i},L{i}>0,T_av-{t2e}-t_tr>0.000001,T_av-{t2e}-2*t_tr-t_ch>=0.25)")
    put(ws, i, 14, f"=IF(M{i},MAX(0,MIN(I{i},T_av-({t2e}+2*t_tr+t_ch)-t_tr)),0)", fmt="0.00")
    put(ws, i, 15, f"=K{i}*1+M{i}*1", fmt="0")
    put(ws, i, 16, f"=(J{i}+L{i}+N{i})/D{i}", fmt="0.000")
    put(ws, i, 17, f"=(E{i}+E_tr*(2+2*O{i})/P{i})/(eta_ch*eta_bat)", fmt="0.00")
    put(ws, i, 18, f"={E}T{i}/B{i}+C{i}*turn_t/3600*Q_turn", fmt="0.00")
    put(ws, i, 19, f"=R{i}/D{i}", fmt="0.00")
    put(ws, i, 20, f"=tank_D/S{i}", fmt="0.00")
    put(ws, i, 21, f"=MAX(0,((T_av-2*t_tr)*S{i}+2*t_tr*Q_tr)/tank_D-1)", fmt="0.00")
    put(ws, i, 22, f"=(T_av-2*t_tr-U{i}*t_refuel_h)/D{i}", fmt="0.000")
    put(ws, i, 23, f"=R{i}+2*t_tr*Q_tr/V{i}", fmt="0.00")
    put(ws, i, 24, f"=T{i}/D{i}", fmt="0.000")
ws.freeze_panes = "B2"

# ============================================================ Экономика
ws = wb.create_sheet("Экономика")
EC = ["Режим", "Труд ЭТ", "Труд Д", "Энергия ЭТ", "Топливо Д", "Ремонт ЭТ", "Ремонт Д", "A_год, га", "Циклов/год",
      "Ресурс батареи, лет", "Замен батареи", "Капит. ЭТ", "Капит. Д", "C_пр ЭТ", "C_пр Д", "C_сов ЭТ", "C_сов Д"]
header(ws, 1, [h + (", тг/га" if j in range(1, 7) or j >= 11 else "") for j, h in enumerate(EC)], [16] + [11] * 16)
for i in range(2, NR + 2):
    s = "Смена!"
    put(ws, i, 1, f"=Смена!A{i}", F_LINK)
    put(ws, i, 2, f"=wage_h*({s}D{i}+2*t_tr*(1+{s}O{i})/{s}P{i})", fmt="#,##0")
    put(ws, i, 3, f"=wage_h*({s}D{i}+2*t_tr/{s}V{i})", fmt="#,##0")
    put(ws, i, 4, f"={s}Q{i}*c_el", fmt="#,##0")
    put(ws, i, 5, f"={s}W{i}*c_D", fmt="#,##0")
    put(ws, i, 6, f"=k_RM_ET*rm_h_D*{s}D{i}", fmt="#,##0")
    put(ws, i, 7, f"=rm_h_D*{s}D{i}", fmt="#,##0")
    put(ws, i, 8, f"=T_year_h/{s}D{i}", fmt="0.0")
    put(ws, i, 9, f"=H{i}*{s}E{i}/E_bat_kWh", fmt="0.0")
    put(ws, i, 10, f"=MIN(bat_cycles/I{i},bat_cal_y)", fmt="0.0")
    put(ws, i, 11, f"=MAX(0,life_y/J{i}-1)", fmt="0.00")
    put(ws, i, 12, f"=(price_ET*(k_CRF-salvage*k_SFF)+C_bat*K{i}/(1+r_disc)^J{i}*k_CRF)/H{i}", fmt="#,##0")
    put(ws, i, 13, f"=price_D*(k_CRF-salvage*k_SFF)/H{i}", fmt="#,##0")
    put(ws, i, 14, f"=B{i}+D{i}+F{i}", fmt="#,##0")
    put(ws, i, 15, f"=C{i}+E{i}+G{i}", fmt="#,##0")
    put(ws, i, 16, f"=N{i}+L{i}", fmt="#,##0")
    put(ws, i, 17, f"=O{i}+M{i}", fmt="#,##0")
ws.freeze_panes = "B2"

# ============================================================ Веса (МАИ)
ws = wb.create_sheet("Веса")
ws.column_dimensions["A"].width = 18
for j in range(2, 14):
    ws.column_dimensions[L(j)].width = 11
ws["A1"] = "Матрица парных сравнений групп (изменяемая: верхний треугольник; нижний — обратные значения)"
ws["A1"].font = F_B
GA = ["Энергет.", "Технолог.", "Экономич.", "Экологич.", "Автономн."]
r = 3
put(ws, r - 1, 1, "Группа", F_H, fill=FILL_H)
for j, g in enumerate(GA, 2):
    put(ws, r - 1, j, g, F_H, fill=FILL_H)
for j, h in enumerate(["Геом. среднее", "Вес w", "(A·w)ᵢ/wᵢ"], 7):
    put(ws, r - 1, j, h, F_H, fill=FILL_H)
for i in range(5):
    put(ws, r + i, 1, GROUPS[i], F_B)
    for j in range(5):
        if j > i:
            put(ws, r + i, 2 + j, AHP_GROUPS[i][j], F_IN, fmt="0.000", fill=FILL_IN)
        elif j == i:
            put(ws, r + i, 2 + j, 1)
        else:
            put(ws, r + i, 2 + j, f"=1/{L(2 + i)}{r + j}", fmt="0.000")
    put(ws, r + i, 7, f"=GEOMEAN(B{r + i}:F{r + i})", fmt="0.0000")
    put(ws, r + i, 8, f"=G{r + i}/SUM($G${r}:$G${r + 4})", fmt="0.0000")
# транспонированная строка весов для SUMPRODUCT
for j in range(5):
    put(ws, r + 6, 2 + j, f"=H{r + j}", fmt="0.0000")
put(ws, r + 6, 1, "wᵀ")
for i in range(5):
    put(ws, r + i, 9, f"=SUMPRODUCT(B{r + i}:F{r + i},$B${r + 6}:$F${r + 6})/H{r + i}", fmt="0.0000")
put(ws, r + 7, 1, "λ_max")
put(ws, r + 7, 2, f"=AVERAGE(I{r}:I{r + 4})", fmt="0.0000")
put(ws, r + 8, 1, "CI")
put(ws, r + 8, 2, f"=(B{r + 7}-5)/4", fmt="0.0000")
put(ws, r + 9, 1, "CR (RI = 1,12)")
put(ws, r + 9, 2, f"=B{r + 8}/1.12", fmt="0.0000", fill=FILL_KEY)
gw = {GROUPS[i]: f"$H${r + i}" for i in range(5)}
# локальные матрицы
r = r + 12
ws.cell(r, 1, "Локальные матрицы (внутри групп)").font = F_B
r += 1
loc = {}
for g in GROUPS:
    Mx = AHP_LOCAL[g]
    codes = [c[0] for c in CRITERIA if c[1] == g]
    n = len(codes)
    put(ws, r, 1, g, F_H, fill=FILL_H)
    for j, c in enumerate(codes, 2):
        put(ws, r, j, c, F_H, fill=FILL_H)
    put(ws, r, 2 + n, "Геом. ср.", F_H, fill=FILL_H)
    put(ws, r, 3 + n, "Лок. вес", F_H, fill=FILL_H)
    for i in range(n):
        rr = r + 1 + i
        put(ws, rr, 1, codes[i], F_B)
        for j in range(n):
            if j > i:
                put(ws, rr, 2 + j, Mx[i][j], F_IN, fmt="0.000", fill=FILL_IN)
            elif j == i:
                put(ws, rr, 2 + j, 1)
            else:
                put(ws, rr, 2 + j, f"=1/{L(2 + i)}{r + 1 + j}", fmt="0.000")
        put(ws, rr, 2 + n, f"=GEOMEAN(B{rr}:{L(1 + n)}{rr})", fmt="0.0000")
        put(ws, rr, 3 + n, f"={L(2 + n)}{rr}/SUM(${L(2 + n)}${r + 1}:${L(2 + n)}${r + n})", fmt="0.0000")
        loc[codes[i]] = f"${L(3 + n)}${rr}"
    r += n + 2
# итоговые веса
r += 1
ws.cell(r, 1, "Итоговые веса").font = F_B
r += 1
for j, h in enumerate(["Код", "Группа", "w МАИ", "w энтроп.", "w итог"], 1):
    put(ws, r, j, h, F_H, fill=FILL_H)
WROW = {}
for i, (code, grp, *_r) in enumerate(CRITERIA):
    rr = r + 1 + i
    put(ws, rr, 1, code, F_B)
    put(ws, rr, 2, grp)
    put(ws, rr, 3, f"={gw[grp]}*{loc[code]}", fmt="0.0000")
    put(ws, rr, 4, f"=Энтропия!{L(2 + i)}$40", F_LINK, fmt="0.0000")
    put(ws, rr, 5, f"=(alpha_w*C{rr}+(1-alpha_w)*D{rr})/(alpha_w*SUM($C${r + 1}:$C${r + 10})+(1-alpha_w)*SUM($D${r + 1}:$D${r + 10}))", fmt="0.0000", fill=FILL_KEY)
    WROW[code] = f"Веса!$E${rr}"

# ============================================================ Показатели
ws = wb.create_sheet("Показатели")
hdr = ["Режим", "Допустим ЭТ и Д", "S_зад, га"] + [f"{c} ЭТ" for c in CODES] + [f"{c} Д" for c in CODES] + [f"d {c}" for c in CODES] + ["Индекс I"]
header(ws, 1, hdr, [16, 10, 9] + [9] * 31)
sign = {c[0]: c[4] for c in CRITERIA}
for i in range(2, NR + 2):
    s, e, en = "Смена!", "Экономика!", "Энергетика!"
    put(ws, i, 1, f"=Смена!A{i}", F_LINK)
    put(ws, i, 2, f"=AND({en}Q{i},{en}R{i})")
    put(ws, i, 3, f"=IF(S_task>0,S_task,{s}V{i})", fmt="0.00")
    ET = {"I1": f"=3.6*{s}Q{i}", "I2": f"=3.6*{s}Q{i}*PEF_el", "I3": f"=MIN(1,{s}P{i}/C{i})", "I4": f"=100*{en}I{i}",
          "I5": f"=100*{en}O{i}", "I6": f"={e}N{i}", "I7": f"={e}P{i}", "I8": f"={s}Q{i}*EF_grid", "I9": "=0",
          "I10": f"=MIN(1,{s}H{i}/C{i})"}
    DD = {"I1": f"={s}W{i}*LHV_D", "I2": f"={s}W{i}*LHV_D*PEF_D", "I3": f"=MIN(1,{s}V{i}/C{i})", "I4": f"=100*{en}K{i}",
          "I5": f"=100*{en}P{i}", "I6": f"={e}O{i}", "I7": f"={e}Q{i}", "I8": f"={s}W{i}*EF_D", "I9": f"={s}W{i}*EF_D",
          "I10": f"=MIN(1,{s}X{i}/C{i})"}
    for j, c in enumerate(CODES):
        put(ws, i, 4 + j, ET[c], fmt="0.000")
        put(ws, i, 14 + j, DD[c], fmt="0.000")
        x, y = f"{L(4 + j)}{i}", f"{L(14 + j)}{i}"
        put(ws, i, 24 + j, f"=IF(MAX(ABS({x}),ABS({y}))=0,0,{sign[c]}*({x}-{y})/MAX(ABS({x}),ABS({y})))", fmt="0.000")
    put(ws, i, 34, "=" + "+".join(f"{WROW[c]}*{L(24 + j)}{i}" for j, c in enumerate(CODES)), fmt="0.000", fill=FILL_KEY)
ws.freeze_panes = "B2"
# диаграмма индекса
ch = BarChart()
ch.type = "bar"
ch.title = "Интегральный индекс I по режимам (I > 0 — преимущество ЭТ)"
ch.y_axis.title = "I"
data = Reference(ws, min_col=34, min_row=1, max_row=NR + 1)
cats = Reference(ws, min_col=1, min_row=2, max_row=NR + 1)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
ch.height, ch.width = 11, 20
ch.legend = None
ws.add_chart(ch, "B20")

# ============================================================ Энтропия
ws = wb.create_sheet("Энтропия")
header(ws, 1, ["Альтернатива", *CODES, "Допустим"], [18] + [10] * 11)
alt = []
for i in range(2, NR + 2):
    alt.append((i, "ЭТ", 4))
    alt.append((i, "Д", 14))
r = 2
for (src, mach, off) in alt:
    put(ws, r, 1, f'=Показатели!A{src}&" {mach}"', F_LINK)
    for j in range(10):
        put(ws, r, 2 + j, f"=Показатели!{L(off + j)}{src}", F_LINK, fmt="0.000")
    put(ws, r, 12, f"=Показатели!B{src}", F_LINK)
    r += 1
last = r - 1  # 33
put(ws, 35, 1, "m (допустимых)", F_B)
put(ws, 35, 2, f"=COUNTIF(L2:L{last},TRUE)")
put(ws, 36, 1, "min", F_B)
put(ws, 37, 1, "max", F_B)
put(ws, 38, 1, "Σ z", F_B)
put(ws, 39, 1, "Энтропия E", F_B)
put(ws, 40, 1, "Энтропийный вес", F_B)
# z и p·ln p — вспомогательные блоки справа
for j in range(10):
    c = L(2 + j)
    put(ws, 36, 2 + j, f"=_xlfn.MINIFS({c}2:{c}{last},$L$2:$L${last},TRUE)", fmt="0.000")
    put(ws, 37, 2 + j, f"=_xlfn.MAXIFS({c}2:{c}{last},$L$2:$L${last},TRUE)", fmt="0.000")
    zc = L(14 + j)
    pc_ = L(25 + j)
    put(ws, 1, 14 + j, f"z {CODES[j]}", F_H, fill=FILL_H)
    put(ws, 1, 25 + j, f"p·ln p {CODES[j]}", F_H, fill=FILL_H)
    for rr in range(2, last + 1):
        put(ws, rr, 14 + j, f"=IF($L{rr},IF({c}$37>{c}$36,({c}{rr}-{c}$36)/({c}$37-{c}$36),0)+0.000001,0)", fmt="0.0000")
        put(ws, rr, 25 + j, f"=IF($L{rr},{zc}{rr}/{zc}$38*LN({zc}{rr}/{zc}$38),0)", fmt="0.0000")
    put(ws, 38, 2 + j, f"=SUM({zc}2:{zc}{last})", fmt="0.0000")
    ws.cell(38, 14 + j, f"=SUM({zc}2:{zc}{last})")
    put(ws, 39, 2 + j, f"=-SUM({pc_}2:{pc_}{last})/LN($B$35)", fmt="0.0000")
for j in range(10):
    put(ws, 40, 2 + j, f"=(1-{L(2 + j)}39)/(10-SUM($B$39:$K$39))", fmt="0.0000", fill=FILL_KEY)
# в формуле p используется Σz в строке 38 столбцов z
for j in range(10):
    zc = L(14 + j)
    for rr in range(2, last + 1):
        ws.cell(rr, 25 + j).value = f"=IF($L{rr},{zc}{rr}/{zc}$38*LN({zc}{rr}/{zc}$38),0)"

# ============================================================ Устойчивость (значения из Python)
ws = wb.create_sheet("Устойчивость")
mc = pd.read_csv(T / "mc_summary.csv")
header(ws, 1, ["Сценарий", "Машина", "a, м", "v, км/ч", "I среднее", "I 5 %", "I 95 %", "P(I>0)", "P(C_сов ЭТ < Д)"], [10, 12, 8, 8, 10, 10, 10, 10, 14])
for i, x in enumerate(mc.itertuples(), 2):
    for j, v in enumerate([x.scen, x.implement, x.a, x.v, x.I_mean, x.I_p05, x.I_p95, x.P_pos, x.P_cost], 1):
        put(ws, i, j, round(v, 4) if isinstance(v, float) else v, F_IN, fmt="0.000" if j >= 5 else None)
c = ws.cell(1, 11, "Значения получены программой src/analysis.py (1000 испытаний Монте-Карло: параметры — равномерно в диапазонах листа «Исходные», веса — распределение Дирихле с параметром 50·w). S1 — полная смена, S2 — задание 0,5 га.")
c.font = F_N
c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["K"].width = 60
src = pd.read_csv(T / "sensitivity_src.csv", index_col=0)
r = len(mc) + 4
header(ws, r, ["Параметр", "SRC S1", "SRC S2"])
for i, (k, x) in enumerate(src.iterrows(), r + 1):
    put(ws, i, 1, k)
    put(ws, i, 2, round(x.S1, 4), F_IN, fmt="0.000")
    put(ws, i, 3, round(x.S2, 4), F_IN, fmt="0.000")

for w_ in wb.worksheets:
    w_.sheet_view.zoomScale = 90
OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print("saved", OUT)

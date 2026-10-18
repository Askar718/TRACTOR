"""Построитель документа Word с оформлением по ГОСТ 7.32-2017.

Возможности: структурные элементы и разделы, нумерация формул/таблиц/рисунков по разделам,
формулы как редактируемые объекты Word (LaTeX → OMML через pandoc), список источников
в порядке ссылок, содержание с номерами страниц (двухпроходная сборка через LibreOffice).
Разметка в тексте: *курсив*, **полужирный**, x_{нижн}, x^{верхн}, [@ключ] — ссылка.
"""
import copy
import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor
from lxml import etree

FONT = "Times New Roman"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
TOKEN = re.compile(r"(\*\*.+?\*\*|\*.+?\*|_\{.+?\}|\^\{.+?\}|\[@[\w,\- ]+\])")


def _latex_to_omml(latex):
    # десятичная запятая: 0{,}1 → единое число «0,1»
    latex = re.sub(r"(\d+)\{,\}(\d+)", r"\\text{\1,\2}", latex)
    with tempfile.TemporaryDirectory() as d:
        md = Path(d) / "e.md"
        md.write_text(f"$${latex}$$\n", encoding="utf-8")
        out = Path(d) / "e.docx"
        subprocess.run(["pandoc", str(md), "-o", str(out)], check=True)
        subprocess.run(["unzip", "-q", "-o", str(out), "word/document.xml", "-d", d], check=True)
        xml = (Path(d) / "word" / "document.xml").read_bytes()
    root = etree.fromstring(xml)
    om = root.find(f".//{{{M_NS}}}oMath")
    if om is None:
        raise ValueError(f"pandoc не преобразовал формулу: {latex}")
    return om


class Report:
    def __init__(self, refs):
        self.d = Document()
        self.refs_db = refs          # ключ → библиографическое описание
        self.ref_order = []          # порядок первого упоминания
        self.sec = 0                 # номер раздела (0 — вне раздела; буква — приложение)
        self.n_eq = self.n_tab = self.n_fig = 0
        self.headings = []           # для содержания: (уровень, текст)
        self.counts = {"fig": 0, "tab": 0}
        self._setup()

    # ------------------------------------------------------------ оформление
    def _setup(self):
        s = self.d.sections[0]
        s.page_height, s.page_width = Mm(297), Mm(210)
        s.left_margin, s.right_margin = Mm(30), Mm(15)
        s.top_margin, s.bottom_margin = Mm(20), Mm(20)
        s.footer_distance = Mm(10)
        s.different_first_page_header_footer = True
        st = self.d.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(14)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        pf = st.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        pf.first_line_indent = Cm(1.25)
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.widow_control = True
        for lvl in (1, 2, 3):
            h = self.d.styles[f"Heading {lvl}"]
            h.font.name = FONT
            h.font.size = Pt(14)
            h.font.bold = True
            h.font.italic = False
            h.font.color.rgb = RGBColor(0, 0, 0)
            rpr = h.element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts")
                rpr.append(rf)
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(a), FONT)
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(a)) is not None:
                    del rf.attrib[qn(a)]
            hp = h.paragraph_format
            hp.first_line_indent = Cm(1.25)
            hp.space_before = Pt(12 if lvl == 1 else 6)
            hp.space_after = Pt(12 if lvl == 1 else 6)
            hp.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            hp.keep_with_next = True
            hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        # нумерация страниц: внизу по центру
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.first_line_indent = Cm(0)
        self._field(fp.add_run(), "PAGE")
        # обновление полей при открытии
        settings = self.d.settings.element
        uf = OxmlElement("w:updateFields")
        uf.set(qn("w:val"), "true")
        settings.append(uf)

    @staticmethod
    def _field(run, code):
        f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = code
        f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
        run._r.append(f1); run._r.append(it); run._r.append(f2)

    # ------------------------------------------------------------ текст
    def cite(self, keys):
        nums = []
        for k in [x.strip() for x in keys.split(",")]:
            if k not in self.refs_db:
                raise KeyError(f"нет источника {k}")
            if k not in self.ref_order:
                self.ref_order.append(k)
            nums.append(self.ref_order.index(k) + 1)
        nums = sorted(nums)
        # сжатие диапазонов: 3, 4, 5 → 3–5
        parts, i = [], 0
        while i < len(nums):
            j = i
            while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
                j += 1
            parts.append(f"{nums[i]}–{nums[j]}" if j - i >= 2 else ", ".join(str(n) for n in nums[i:j + 1]))
            i = j + 1
        return "[" + ", ".join(parts) + "]"

    def _runs(self, par, text, size=None, bold=False, italic=False):
        for tok in TOKEN.split(text):
            if not tok:
                continue
            b, it, sub, sup = bold, italic, False, False
            if tok.startswith("**") and tok.endswith("**"):
                tok, b = tok[2:-2], True
            elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
                tok, it = tok[1:-1], True
            elif tok.startswith("_{"):
                tok, sub = tok[2:-1], True
            elif tok.startswith("^{"):
                tok, sup = tok[2:-1], True
            elif tok.startswith("[@"):
                tok = self.cite(tok[2:-1])
            # вложенный курсив внутри индекса не поддерживается — не требуется
            r = par.add_run(tok)
            r.bold, r.italic = b, it
            r.font.subscript, r.font.superscript = sub, sup
            if size:
                r.font.size = Pt(size)
        return par

    def p(self, text, indent=True, align=None, size=None, bold=False, italic=False, keep=False, space_after=0):
        par = self.d.add_paragraph()
        pf = par.paragraph_format
        if not indent:
            pf.first_line_indent = Cm(0)
        if align == "center":
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "left":
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif align == "right":
            pf.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if keep:
            pf.keep_with_next = True
        if space_after:
            pf.space_after = Pt(space_after)
        return self._runs(par, text, size=size, bold=bold, italic=italic)

    def where(self, items):
        """Расшифровка обозначений после формулы: «где X — ...»."""
        for i, (sym, txt) in enumerate(items):
            par = self.d.add_paragraph()
            pf = par.paragraph_format
            pf.first_line_indent = Cm(0)
            pf.left_indent = Cm(0) if i == 0 else Cm(1.0)
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            self._runs(par, ("где " if i == 0 else "") + sym + " – " + txt + (";" if i < len(items) - 1 else "."))

    def bullets(self, items, kind="dash"):
        for i, t in enumerate(items):
            par = self.d.add_paragraph()
            pf = par.paragraph_format
            pf.first_line_indent = Cm(1.25)
            mark = "– " if kind == "dash" else f"{'абвгдежиклмнпрстуфхцчшщэюя'[i]}) " if kind == "alpha" else f"{i + 1}) "
            end = ";" if i < len(items) - 1 else "."
            t = t.rstrip(".;")
            self._runs(par, mark + t + end)

    def pagebreak(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------ заголовки
    def struct(self, title, newpage=True, toc=True):
        """Структурный элемент: РЕФЕРАТ, ВВЕДЕНИЕ, ЗАКЛЮЧЕНИЕ и т.п."""
        if newpage:
            self.pagebreak()
        h = self.d.add_paragraph(style="Heading 1")
        h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h.paragraph_format.first_line_indent = Cm(0)
        h.add_run(title.upper())
        if toc:
            self.headings.append((1, title.upper()))
        self.sec = 0
        return h

    def h1(self, num, title):
        self.pagebreak()
        self.sec = num
        self.n_eq = self.n_tab = self.n_fig = 0
        h = self.d.add_paragraph(style="Heading 1")
        h.add_run(f"{num} {title}")
        self.headings.append((1, f"{num} {title}"))

    def h2(self, num, title):
        h = self.d.add_paragraph(style="Heading 2")
        h.add_run(f"{num} {title}")
        self.headings.append((2, f"{num} {title}"))

    def h3(self, num, title):
        h = self.d.add_paragraph(style="Heading 3")
        h.add_run(f"{num} {title}")

    def appendix(self, letter, kind, title):
        self.pagebreak()
        self.sec = letter
        self.n_eq = self.n_tab = self.n_fig = 0
        h = self.d.add_paragraph(style="Heading 1")
        h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h.paragraph_format.first_line_indent = Cm(0)
        h.paragraph_format.space_after = Pt(0)
        h.add_run(f"ПРИЛОЖЕНИЕ {letter}")
        self.p(f"({kind})", indent=False, align="center")
        self.p(title, indent=False, align="center", bold=True, space_after=12)
        self.headings.append((1, f"Приложение {letter} ({kind}) {title}"))

    def _num(self, n):
        return f"{self.sec}.{n}" if self.sec else f"{n}"

    # ------------------------------------------------------------ формулы
    def eq(self, latex, label=None):
        self.n_eq += 1
        num = self._num(self.n_eq)
        par = self.d.add_paragraph()
        pf = par.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.space_before = Pt(6)
        pf.space_after = Pt(6)
        pf.keep_with_next = True
        ts = pf.tab_stops
        ts.add_tab_stop(Cm(8.25), WD_TAB_ALIGNMENT.CENTER)
        ts.add_tab_stop(Cm(16.5), WD_TAB_ALIGNMENT.RIGHT)
        par.add_run("\t")
        om = _latex_to_omml(latex)
        par._p.append(parse_xml(etree.tostring(om)))
        par.add_run(f"\t({num})")
        return f"({num})"

    # ------------------------------------------------------------ таблицы
    def table(self, caption, header, rows, widths=None, size=12, align=None, header_rows=1, continued=False):
        self.n_tab += 1
        self.counts["tab"] += 1
        num = self._num(self.n_tab)
        cap = self.d.add_paragraph()
        cap.paragraph_format.first_line_indent = Cm(0)
        cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        cap.paragraph_format.keep_with_next = True
        cap.paragraph_format.space_before = Pt(6)
        self._runs(cap, f"Таблица {num} – {caption}")
        hdr = header if isinstance(header[0], (list, tuple)) else [header]
        ncol = len(hdr[-1])
        t = self.d.add_table(rows=0, cols=ncol)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        total = 16.5
        if widths is None:
            widths = [total / ncol] * ncol
        else:
            k = total / sum(widths)
            widths = [w * k for w in widths]
        for hi, h in enumerate(hdr):
            cells = t.add_row().cells
            j = 0
            for item in h:
                span = 1
                if isinstance(item, tuple):
                    item, span = item
                c = cells[j]
                if span > 1:
                    c = c.merge(cells[j + span - 1])
                self._cell(c, str(item), size, bold=False, center=True)
                j += span
            self._repeat_header(t.rows[-1])
        for r in rows:
            cells = t.add_row().cells
            for j, v in enumerate(r):
                a = (align[j] if align else ("left" if j == 0 else "center"))
                self._cell(cells[j], str(v), size, center=(a == "center"))
        for row in t.rows:
            for j, c in enumerate(row.cells):
                c.width = Cm(widths[j])
        self._autofit_off(t, widths)
        sp = self.d.add_paragraph()
        sp.paragraph_format.first_line_indent = Cm(0)
        sp.paragraph_format.line_spacing = 1.0
        return f"{num}"

    def _cell(self, c, text, size, bold=False, center=True):
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        par = c.paragraphs[0]
        pf = par.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        lines = text.split("\n")
        for i, line in enumerate(lines):
            if i:
                par = c.add_paragraph()
                par.paragraph_format.first_line_indent = Cm(0)
                par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
            self._runs(par, line, size=size, bold=bold)

    @staticmethod
    def _repeat_header(row):
        trPr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        trPr.append(el)

    @staticmethod
    def _autofit_off(t, widths):
        tblPr = t._tbl.tblPr
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tblPr.append(lay)
        grid = t._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(w / 2.54 * 1440)))

    # ------------------------------------------------------------ рисунки
    def fig(self, path, caption, width=None):
        if width is None:   # натуральная величина рисунка, построенного с разрешением 300 точек на дюйм
            from PIL import Image
            with Image.open(path) as im:
                width = min(16.5, im.width / 300 * 2.54)
        self.n_fig += 1
        self.counts["fig"] += 1
        num = self._num(self.n_fig)
        par = self.d.add_paragraph()
        pf = par.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.keep_with_next = True
        pf.line_spacing = 1.0
        pf.space_before = Pt(6)
        par.add_run().add_picture(str(path), width=Cm(width))
        cap = self.d.add_paragraph()
        cap.paragraph_format.first_line_indent = Cm(0)
        cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(6)
        self._runs(cap, f"Рисунок {num} – {caption}")
        return num

    def note(self, text):
        """Примечание к таблице (ГОСТ 7.32-2017, 6.6.8): с абзацного отступа, после таблицы."""
        par = self.d.add_paragraph()
        pf = par.paragraph_format
        pf.first_line_indent = Cm(1.25)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.line_spacing = 1.0
        pf.space_before = Pt(3)
        pf.space_after = Pt(6)
        self._runs(par, f"Примечание – {text}", size=12)
        return par

    # ------------------------------------------------------------ содержание
    def toc_placeholder(self):
        self._toc_anchor = self.d.add_paragraph()
        self._toc_anchor.paragraph_format.first_line_indent = Cm(0)

    def fill_toc(self, pages):
        """pages: номера страниц для self.headings (в том же порядке)."""
        from docx.enum.text import WD_TAB_LEADER
        for (lvl, text), pg in zip(self.headings, pages):
            par = self._toc_anchor.insert_paragraph_before()
            pf = par.paragraph_format
            pf.first_line_indent = Cm(0)
            pf.left_indent = Cm(0 if lvl == 1 else 0.6)
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.tab_stops.add_tab_stop(Cm(16.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            par.add_run(text)
            par.add_run(f"\t{pg if pg else ''}")

    # ------------------------------------------------------------ источники
    def references(self):
        self.struct("Список использованных источников")
        for i, k in enumerate(self.ref_order, 1):
            par = self.d.add_paragraph()
            par.paragraph_format.first_line_indent = Cm(1.25)
            self._runs(par, f"{i}. {self.refs_db[k]}")

    def save(self, path):
        self.d.save(path)
